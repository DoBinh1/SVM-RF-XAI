import sys
from pptx import Presentation

def main():
    hust_slides_path = r"d:\[Lab] HUST\nhà máy\slides\BaiGiang_CWRU_HUST_v2.pptx"
    try:
        prs = Presentation(hust_slides_path)
    except Exception as e:
        print(f"Error loading HUST slide file: {e}")
        sys.exit(1)
        
    output_path = r"d:\[Lab] HUST\nhà máy\verify_v2_report.txt"
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(f"Total slides in generated file: {len(prs.slides)}\n")
        
        slide = prs.slides[11]
        title_text = slide.shapes.title.text if slide.shapes.title else ''
        f.write(f"Slide 11 Title: '{title_text}'\n")
        
        f.write("Slide 11 Shapes Content:\n")
        for idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                txt = shape.text.strip().replace('\n', ' ')
                f.write(f"  Shape {idx} (name: '{shape.name}'): '{txt}'\n")

if __name__ == "__main__":
    main()
